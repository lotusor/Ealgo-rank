"""「导出学校赛情」的接口与口径测试。

覆盖三类断言，缺任何一类都可能是「逻辑对但通道错」：

1. **权限**：角色矩阵、跨校、无校、降级、转校、他人任务 id。
2. **口径**：参赛/异常次数、平台 rating 的赛季末值、总rating 取整与并列排名、
   姓名回退不抖 UUID、学号文本格式。
3. **安全**：公式注入、X-Accel-Redirect、过期 410、幂等指纹含 actor、
   查询数不随人数线性增长（防 N+1 回归）。

注意：CI 的测试库是 **SQLite**。任何 PostgreSQL 专有语法（`DISTINCT ON`）
都会在这里直接炸——所以「本文件能跑通」本身就是 CI 红线的守门人。
"""
import os
import shutil
import tempfile
import unittest
import zipfile
from datetime import timedelta
from pathlib import Path

from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from apps.accounts.models import PlatformAccount, User, UserRole
from apps.common.models import ExcludeReason, Platform
from apps.contests.models import Contest, Participation
from apps.exports import services
from apps.exports.models import ExportAuditLog, ExportEvent, ExportJob
from apps.exports.tasks import run_export
from apps.ranking.models import RankSnapshot
from apps.schools.models import School

BASE = "/api/v1"
EXPORT_LIST = f"{BASE}/exports/"
YEAR = 2026


def export_detail(pk):
    return f"{BASE}/exports/{pk}/"


def export_download(pk):
    return f"{BASE}/exports/{pk}/download/"


def make_user(username, role=UserRole.USER, school=None, **extra):
    return User.objects.create_user(
        username=username, password="Test1234!", role=role, school=school,
        **extra)


class ExportTestBase(APITestCase):
    """把导出目录指到临时目录，避免污染仓库。"""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._tmpdir = tempfile.mkdtemp(prefix="rank-export-test-")
        cls._override = override_settings(EXPORT_ROOT=Path(cls._tmpdir))
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._tmpdir, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        # 限流用 LocMemCache，进程内跨测试共享；不清会让同 pk 的账号在后续
        # 用例里一上来就吃 429（与本功能无关的夹具污染）。
        from django.core.cache import cache

        cache.clear()

        self.school_a = School.objects.create(name="甲大学", code="jia",
                                              short_name="甲大")
        self.school_b = School.objects.create(name="乙大学", code="yi",
                                              short_name="乙大")

        self.boss = make_user("root", UserRole.SUPER_ADMIN)
        self.admin_a = make_user("admin_a", UserRole.SCHOOL_ADMIN,
                                 school=self.school_a)
        self.admin_b = make_user("admin_b", UserRole.SCHOOL_ADMIN,
                                 school=self.school_b)
        self.homeless_admin = make_user("admin_none", UserRole.SCHOOL_ADMIN)
        self.plain = make_user("student_plain", UserRole.USER,
                               school=self.school_a)

        self.c1 = Contest.objects.create(
            platform=Platform.CODEFORCES, external_id="c1", name="春季赛",
            is_rated=True, participant_count=500,
            start_time=timezone.make_aware(
                timezone.datetime(YEAR, 3, 1, 12, 0)))
        self.c2 = Contest.objects.create(
            platform=Platform.CODEFORCES, external_id="c2", name="夏季赛",
            is_rated=True, participant_count=500,
            start_time=timezone.make_aware(
                timezone.datetime(YEAR, 6, 1, 12, 0)))
        self.c_old = Contest.objects.create(
            platform=Platform.CODEFORCES, external_id="c0", name="去年赛",
            is_rated=True, participant_count=500,
            start_time=timezone.make_aware(
                timezone.datetime(YEAR - 1, 11, 1, 12, 0)))

        # 学生甲：有真名与学号，两场 2026 比赛（其中一场作弊）
        self.s1 = make_user("s1", school=self.school_a, real_name="张一",
                            student_no="2509030326")
        self.pa_s1 = PlatformAccount.objects.create(
            user=self.s1, platform=Platform.CODEFORCES, handle="cf_zhang",
            school=self.school_a)
        Participation.objects.create(
            contest=self.c1, platform_account=self.pa_s1, handle="cf_zhang",
            rank=5, new_rating=1300)
        Participation.objects.create(
            contest=self.c2, platform_account=self.pa_s1, handle="cf_zhang",
            rank=9, new_rating=1400,
            is_excluded=True, exclude_reason=ExcludeReason.CHEATER)
        Participation.objects.create(
            contest=self.c_old, platform_account=self.pa_s1,
            handle="cf_zhang", rank=3, new_rating=1250)

        # 学生乙：无学号，牛客 rating；与甲同分 → 并列名次
        self.s2 = make_user("s2", school=self.school_a, real_name="李二")
        self.pa_s2 = PlatformAccount.objects.create(
            user=self.s2, platform=Platform.NOWCODER, handle="nc_li",
            school=self.school_a)
        Participation.objects.create(
            contest=self.c1, platform_account=self.pa_s2, handle="nc_li",
            rank=20, new_rating=1200)

        # 学生丙：本校本季最高分
        self.s3 = make_user("s3", school=self.school_a, real_name="王三",
                            student_no="2509030328")
        self.pa_s3 = PlatformAccount.objects.create(
            user=self.s3, platform=Platform.ATCODER, handle="at_wang",
            school=self.school_a)
        Participation.objects.create(
            contest=self.c1, platform_account=self.pa_s3, handle="at_wang",
            rank=1, new_rating=900)

        # 学生丁：最低分
        self.s4 = make_user("s4", school=self.school_a, real_name="赵四",
                            student_no="2509030329")
        self.pa_s4 = PlatformAccount.objects.create(
            user=self.s4, platform=Platform.CODEFORCES, handle="cf_zhao",
            school=self.school_a)
        Participation.objects.create(
            contest=self.c1, platform_account=self.pa_s4, handle="cf_zhao",
            rank=99, new_rating=1000)

        # 学生戊：passport 首登占位名（UUID），且无参赛 → 默认不出现
        self.s5 = make_user("0f8b5c1e-1111-4a2b-9c3d-eeeeeeeeeeee",
                            school=self.school_a,
                            passport_user_id="0f8b5c1e-1111-4a2b-9c3d-eeeeeeeeeeee")

        computed = timezone.now()
        for user, score in ((self.s3, 1700.0), (self.s1, 1460.4),
                            (self.s2, 1460.0), (self.s4, 1000.0)):
            RankSnapshot.objects.create(
                scope=RankSnapshot.Scope.STUDENT, period=str(YEAR), user=user,
                school=user.school, rank=1, total_score=score,
                contest_count=1, computed_at=computed)

    # ------------------------------------------------------------------ 工具
    def close(self, resp):
        """消费并关闭响应。

        开发配置走 Django 直出（`FileResponse`），文件句柄在整个响应期间都开着；
        测试里不消费流就等于不关文件。Windows 上后续对**同名文件**做 `os.replace`
        会直接 `WinError 5 拒绝访问`（job 主键在事务回滚后会复用，文件名相同）。
        生产走 `X-Accel-Redirect`，Python 不持句柄，不存在这个问题。
        """
        stream = getattr(resp, "streaming_content", None)
        if stream is not None:
            for _ in stream:
                pass
        resp.close()
        return resp

    def fetch_download(self, job):
        """下载并立即关闭响应；绝大多数断言只看状态码与响应头。"""
        return self.close(self.client.get(export_download(job)))

    def payload(self, **over):
        data = {
            "scope_kind": "school",
            "school_ids": [self.school_a.id],
            "year": YEAR,
            "layout": "per_school",
        }
        data.update(over)
        return data

    def create_export(self, actor, **over):
        self.client.force_authenticate(actor)
        return self.client.post(EXPORT_LIST, self.payload(**over),
                                format="json")

    def sheet_rows(self, job):
        """读回生成的 xlsx（第一张业务表）。"""
        path = Path(self._tmpdir) / job.file_path
        wb = load_workbook(path)
        return wb, wb.worksheets[0]


class ExportPermissionTests(ExportTestBase):
    def test_plain_user_forbidden(self):
        resp = self.create_export(self.plain)
        self.assertEqual(resp.status_code, 403)

    def test_admin_without_school_forbidden(self):
        resp = self.create_export(self.homeless_admin)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data["code"], "no_school")

    def test_admin_cannot_request_other_school(self):
        resp = self.create_export(self.admin_a,
                                  school_ids=[self.school_b.id])
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data["code"], "school_mismatch")
        # 越权尝试必须留痕，否则探测行为无从发现
        self.assertTrue(ExportAuditLog.objects.filter(
            event=ExportEvent.DENIED, reason="school_mismatch").exists())

    def test_admin_cannot_request_all_scope(self):
        resp = self.create_export(self.admin_a, scope_kind="all",
                                  school_ids=[])
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data["code"], "scope_forbidden")

    def test_super_admin_can_export_all(self):
        resp = self.create_export(self.boss, scope_kind="all", school_ids=[])
        self.assertEqual(resp.status_code, 201)
        job = ExportJob.objects.get(pk=resp.data["id"])
        # 种子迁移会先建若干学校，所以按「全部启用学校」比对而不是写死两所
        self.assertIn(self.school_a.id, job.scope_school_ids)
        self.assertIn(self.school_b.id, job.scope_school_ids)
        self.assertEqual(len(job.scope_school_ids),
                         School.objects.filter(is_active=True).count())

    def test_admin_only_sees_own_jobs(self):
        other = ExportJob.objects.create(
            actor=self.admin_b, scope_kind="school",
            scope_school_ids=[self.school_b.id], year=YEAR)
        self.client.force_authenticate(self.admin_a)
        resp = self.client.get(export_detail(other.pk))
        self.assertEqual(resp.status_code, 404)

    def test_demoted_admin_cannot_download(self):
        resp = self.create_export(self.admin_a)
        job_id = resp.data["id"]
        self.admin_a.role = UserRole.USER
        self.admin_a.save(update_fields=["role"])

        self.client.force_authenticate(self.admin_a)
        resp = self.fetch_download(job_id)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data["code"], "permission_revoked")

    def test_transferred_admin_cannot_download(self):
        resp = self.create_export(self.admin_a)
        job_id = resp.data["id"]
        self.admin_a.school = self.school_b
        self.admin_a.save(update_fields=["school"])

        self.client.force_authenticate(self.admin_a)
        resp = self.fetch_download(job_id)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data["code"], "school_changed")

    def test_super_admin_can_download_others_job(self):
        resp = self.create_export(self.admin_a)
        job_id = resp.data["id"]
        self.client.force_authenticate(self.boss)
        resp = self.fetch_download(job_id)
        self.assertEqual(resp.status_code, 200)


class ExportContentTests(ExportTestBase):
    def test_header_matches_template(self):
        resp = self.create_export(self.admin_a)
        self.assertEqual(resp.status_code, 201)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)

        self.assertEqual(
            [ws.cell(row=2, column=i).value for i in range(1, 10)],
            ["姓名", "学号", "参赛次数", "异常次数", "牛客rating", "CF rating",
             "Atcoder rating", "总rating", "校排"])
        # 标题合并必须覆盖全部列（模板只并到 H，属瑕疵，这里修正）
        self.assertIn("A1:I1", [str(r) for r in ws.merged_cells.ranges])
        self.assertIn("甲大", ws["A1"].value)
        self.assertTrue(ws["A1"].value.startswith(str(YEAR)))

    def test_counts_and_platform_rating_use_season_window(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)

        rows = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}
        zhang = rows["张一"]

        # 参赛次数只算本赛季（2026）的两场，去年的 c_old 不计
        self.assertEqual(ws.cell(row=zhang, column=3).value, 2)
        # 异常次数只统计 cheater
        self.assertEqual(ws.cell(row=zhang, column=4).value, 1)
        # 平台 rating 取赛季内最后一场（c2 的 1400），不是第一场 1300
        self.assertEqual(ws.cell(row=zhang, column=6).value, 1400)
        # 牛客/AtCoder 未参加 → 留空，不是 0
        self.assertIsNone(ws.cell(row=zhang, column=5).value)
        self.assertIsNone(ws.cell(row=zhang, column=7).value)

    def test_rank_ties_share_position_and_skip(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)

        rank_of = {
            ws.cell(row=r, column=1).value: ws.cell(row=r, column=9).value
            for r in range(3, ws.max_row + 1)
        }
        # 王三 1700 → 1；张一 1460.4 与李二 1460.0 取整后同为 1460 → 并列 2；
        # 赵四 1000 → 4（并列占位后跳号）
        self.assertEqual(rank_of["王三"], 1)
        self.assertEqual(rank_of["张一"], 2)
        self.assertEqual(rank_of["李二"], 2)
        self.assertEqual(rank_of["赵四"], 4)

    def test_total_rating_is_site_score_not_platform_sum(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)

        rows = {ws.cell(row=r, column=1).value: r for r in range(3, ws.max_row + 1)}
        zhang = rows["张一"]
        # 平台只有 CF 1400，但总rating 取站内有效分 1460（取整）
        self.assertEqual(ws.cell(row=zhang, column=8).value, 1460)

    def test_student_no_is_text_formatted(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)
        cell = None
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == "2509030326":
                cell = ws.cell(row=r, column=2)
                break
        self.assertIsNotNone(cell)
        self.assertEqual(cell.number_format, "@")
        self.assertEqual(cell.data_type, "s")

    def test_passport_uuid_not_leaked_and_zero_participation_skipped(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)
        names = [ws.cell(row=r, column=1).value
                 for r in range(3, ws.max_row + 1)]
        self.assertNotIn("0f8b5c1e-1111-4a2b-9c3d-eeeeeeeeeeee", names)
        self.assertNotIn(self.s5.username, names)

    def test_include_zero_participation_lists_unclaimed_student(self):
        resp = self.create_export(self.admin_a,
                                  include_zero_participation=True)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)
        names = [ws.cell(row=r, column=1).value
                 for r in range(3, ws.max_row + 1)]
        # 出现但姓名为空——UUID 占位名不是人名，不能抖出去。
        # school_a 的在用账号：s1–s5 + plain + admin_a = 7
        self.assertIn(None, names)
        self.assertEqual(len(names), 7)

    def test_merged_layout_has_school_and_global_rank(self):
        resp = self.create_export(self.boss, scope_kind="all", school_ids=[],
                                  layout="merged")
        self.assertEqual(resp.status_code, 201)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)
        self.assertEqual(
            [ws.cell(row=2, column=i).value for i in range(1, 12)],
            ["学校", "姓名", "学号", "参赛次数", "异常次数", "牛客rating",
             "CF rating", "Atcoder rating", "总rating", "校排", "全站排名"])

    def test_notes_sheet_carries_caliber_statements(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        wb = load_workbook(Path(self._tmpdir) / job.file_path)
        self.assertIn("说明", wb.sheetnames)
        text = "\n".join(str(c.value) for row in wb["说明"].iter_rows()
                         for c in row if c.value)
        self.assertIn("不等于计分场次", text)
        self.assertIn("不等于三平台 rating 之和", text)

    def test_previous_season_export_uses_old_contests(self):
        resp = self.create_export(self.admin_a, year=YEAR - 1)
        self.assertEqual(resp.status_code, 201)
        job = ExportJob.objects.get(pk=resp.data["id"])
        _wb, ws = self.sheet_rows(job)
        names = [ws.cell(row=r, column=1).value
                 for r in range(3, ws.max_row + 1)]
        self.assertEqual(names, ["张一"])


class ExportSecurityTests(ExportTestBase):
    def test_formula_injection_is_written_as_text(self):
        evil = make_user("evil", school=self.school_a,
                         real_name="=cmd|' /c calc'!A0",
                         student_no="=1+1")
        pa = PlatformAccount.objects.create(
            user=evil, platform=Platform.CODEFORCES, handle="cf_evil",
            school=self.school_a)
        Participation.objects.create(
            contest=self.c1, platform_account=pa, handle="cf_evil", rank=50)

        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        path = Path(self._tmpdir) / job.file_path
        wb = load_workbook(path)
        ws = wb.worksheets[0]

        target = None
        for r in range(3, ws.max_row + 1):
            if ws.cell(row=r, column=2).value == "=1+1":
                target = ws.cell(row=r, column=2)
        self.assertIsNotNone(target)
        self.assertEqual(target.data_type, "s")

        # 原始 XML 里不能出现公式节点 <f>
        with zipfile.ZipFile(path) as zf:
            sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertNotIn("<f>", sheet_xml)

    @override_settings(EXPORT_ACCEL_ENABLED=True)
    def test_download_returns_accel_header(self):
        resp = self.create_export(self.admin_a)
        self.assertEqual(resp.status_code, 201, resp.data)
        job_id = resp.data["id"]
        resp = self.fetch_download(job_id)
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp["X-Accel-Redirect"].startswith(
            "/_files/exports/"))
        self.assertIn("filename*=UTF-8''", resp["Content-Disposition"])
        self.assertEqual(resp["Cache-Control"], "no-store")

    @override_settings(EXPORT_ACCEL_ENABLED=False)
    def test_download_streams_file_when_accel_disabled(self):
        """本地开发无 nginx：必须能直接回流文件，否则下载是 0 字节空响应。

        这条刻意不用 `fetch_download`：它要读的是原始字节流，读完还要显式 close
        ——句柄不关会顶掉后续用例的同名文件写入（见 base 里的 `close` 说明）。
        """
        resp = self.create_export(self.admin_a)
        job_id = resp.data["id"]
        resp = self.client.get(export_download(job_id))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertNotIn("X-Accel-Redirect", resp)
        body = b"".join(resp.streaming_content)
        resp.close()
        self.assertTrue(body.startswith(b"PK"))  # xlsx 就是 zip

    def test_download_records_count_and_audit(self):
        resp = self.create_export(self.admin_a)
        job_id = resp.data["id"]
        self.fetch_download(job_id)
        job = ExportJob.objects.get(pk=job_id)
        self.assertEqual(job.download_count, 1)
        self.assertTrue(ExportAuditLog.objects.filter(
            event=ExportEvent.DOWNLOAD, job_id_snapshot=job_id).exists())

    def test_expired_file_returns_410(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        ExportJob.objects.filter(pk=job.pk).update(
            expires_at=timezone.now() - timedelta(minutes=1))
        resp = self.fetch_download(job.pk)
        self.assertEqual(resp.status_code, 410)
        self.assertEqual(resp.data["code"], "export_expired")

    def test_missing_file_returns_410(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        (Path(self._tmpdir) / job.file_path).unlink()
        resp = self.fetch_download(job.pk)
        self.assertEqual(resp.status_code, 410)

    def test_revoke_deletes_file(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        path = Path(self._tmpdir) / job.file_path
        self.assertTrue(path.exists())

        resp = self.client.delete(export_detail(job.pk))
        self.assertEqual(resp.status_code, 204)
        self.assertFalse(path.exists())
        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.Status.REVOKED)
        self.assertEqual(job.file_path, "")

    def test_same_params_reuse_previous_result(self):
        first = self.create_export(self.admin_a)
        second = self.create_export(self.admin_a)
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 200)
        self.assertEqual(first.data["id"], second.data["id"])

    def test_fingerprint_includes_actor(self):
        """超管与校管同参数不能互相命中对方的任务。"""
        boss_resp = self.create_export(
            self.boss, school_ids=[self.school_a.id])
        admin_resp = self.create_export(self.admin_a)
        self.assertEqual(boss_resp.status_code, 201)
        self.assertEqual(admin_resp.status_code, 201)
        self.assertNotEqual(boss_resp.data["id"], admin_resp.data["id"])

    def test_concurrent_active_job_returns_409(self):
        ExportJob.objects.create(
            actor=self.admin_a, scope_kind="school",
            scope_school_ids=[self.school_a.id], year=YEAR,
            status=ExportJob.Status.PENDING)
        resp = self.create_export(self.admin_a, year=YEAR - 1)
        self.assertEqual(resp.status_code, 409)
        self.assertEqual(resp.data["code"], "export_in_progress")

    def test_year_out_of_range_rejected(self):
        resp = self.create_export(self.admin_a, year=1999)
        self.assertEqual(resp.status_code, 400)

    def test_year_far_future_rejected(self):
        resp = self.create_export(self.admin_a, year=999999)
        self.assertEqual(resp.status_code, 400)

    def test_too_many_schools_rejected(self):
        resp = self.create_export(self.boss, scope_kind="multi",
                                  school_ids=list(range(1, 202)))
        self.assertEqual(resp.status_code, 400)

    def test_unknown_school_rejected(self):
        resp = self.create_export(self.boss, scope_kind="multi",
                                  school_ids=[999999])
        self.assertEqual(resp.status_code, 400)

    def test_empty_scope_rejected(self):
        # 乙校没有任何学生
        resp = self.create_export(self.admin_b,
                                  school_ids=[self.school_b.id])
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data["code"], "empty_scope")

    def test_task_rechecks_permission_before_generating(self):
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        # 模拟「创建后立刻撤职、任务才轮到执行」
        ExportJob.objects.filter(pk=job.pk).update(
            status=ExportJob.Status.PENDING, file_path="")
        self.admin_a.role = UserRole.USER
        self.admin_a.save(update_fields=["role"])

        result = run_export(job.pk)
        self.assertFalse(result["ok"])
        job.refresh_from_db()
        self.assertEqual(job.status, ExportJob.Status.FAILED)
        self.assertIn("permission_revoked", job.error_message)
        self.assertTrue(ExportAuditLog.objects.filter(
            event=ExportEvent.DENIED, reason="permission_revoked",
            job_id_snapshot=job.pk).exists())

    @unittest.skipIf(os.name == "nt", "Windows 无 POSIX 权限位")
    def test_file_permissions_are_world_readable(self):
        """backend 以 root 写、nginx 以 nginx 读，权限错会 403。"""
        resp = self.create_export(self.admin_a)
        job = ExportJob.objects.get(pk=resp.data["id"])
        mode = (Path(self._tmpdir) / job.file_path).stat().st_mode & 0o777
        self.assertEqual(mode, 0o644)


class ExportQueryTests(ExportTestBase):
    def test_query_count_does_not_grow_with_students(self):
        """N+1 回归防线：加人不加查询。

        这是本模块最容易被写坏的地方——`accounts.serializers.get_platform_ratings`
        就是按账号逐个 `.first()` 的反面教材。断言「两条查询数相等」比断言
        「查询数小于 N」更能抓住回归。
        """
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        def measure():
            with CaptureQueriesContext(connection) as ctx:
                services.collect([self.school_a.id], YEAR, {})
            return len(ctx.captured_queries)

        before = measure()

        for i in range(10):
            student = make_user(f"bulk{i}", school=self.school_a,
                                real_name=f"批量{i}")
            pa = PlatformAccount.objects.create(
                user=student, platform=Platform.CODEFORCES,
                handle=f"bulk_{i}", school=self.school_a)
            Participation.objects.create(
                contest=self.c1, platform_account=pa, handle=f"bulk_{i}",
                rank=i + 10, new_rating=900 + i)

        after = measure()
        self.assertEqual(before, after,
                         "导出取数的查询数随人数增长，说明出现了 N+1")


class ExportModelTests(ExportTestBase):
    def test_partial_unique_index_blocks_second_active_job(self):
        from django.db import IntegrityError, transaction

        ExportJob.objects.create(actor=self.admin_a, scope_kind="school",
                                 scope_school_ids=[self.school_a.id],
                                 year=YEAR)
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                ExportJob.objects.create(actor=self.admin_a,
                                         scope_kind="school",
                                         scope_school_ids=[self.school_a.id],
                                         year=YEAR)

    def test_finished_job_allows_new_one(self):
        job = ExportJob.objects.create(
            actor=self.admin_a, scope_kind="school",
            scope_school_ids=[self.school_a.id], year=YEAR)
        job.status = ExportJob.Status.SUCCEEDED
        job.save(update_fields=["status"])
        # 不抛异常即通过
        ExportJob.objects.create(actor=self.admin_a, scope_kind="school",
                                 scope_school_ids=[self.school_a.id],
                                 year=YEAR)
